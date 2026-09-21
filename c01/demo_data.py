from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def _write_supplier_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    rows = _read_csv(csv_path)
    if not rows:
        raise ValueError(f"Demo source CSV is empty: {csv_path}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Catalogue"

    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Supplier SKU columns should normally be text. One intentional demo row is
    # stored as the numeric value 721 with the Excel display format 000000 to
    # demonstrate the leading-zero risk that cannot be safely guessed.
    for row_idx in range(2, ws.max_row + 1):
        sku_cell = ws.cell(row=row_idx, column=1)
        if str(sku_cell.value).strip() == "721":
            sku_cell.value = 721
            sku_cell.number_format = "000000"
        elif sku_cell.value not in (None, ""):
            sku_cell.number_format = "@"

    widths = {
        1: 18,
        2: 36,
        3: 18,
        4: 14,
        5: 12,
        6: 20,
        7: 28,
        8: 18,
        9: 28,
        10: 14,
        11: 12,
        12: 10,
        13: 10,
        14: 12,
        15: 10,
    }
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def generate_demo_files(source_dir: str | Path, output_dir: str | Path | None = None) -> tuple[Path, Path]:
    source_dir = Path(source_dir)
    output_dir = Path(output_dir) if output_dir is not None else source_dir
    previous = output_dir / "supplier_previous.xlsx"
    current = output_dir / "supplier_current.xlsx"
    _write_supplier_xlsx(source_dir / "source_previous.csv", previous)
    _write_supplier_xlsx(source_dir / "source_current.csv", current)
    return previous, current
