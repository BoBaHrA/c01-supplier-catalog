from __future__ import annotations

import csv
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CanonicalProduct, Change, NormalizationRecord, ValidationResult


CSV_HEADERS = [
    "Type",
    "SKU",
    "GTIN, UPC, EAN, or ISBN",
    "Name",
    "Published",
    "Tax status",
    "Tax class",
    "In stock?",
    "Stock",
    "Regular price",
    "Categories",
    "Parent",
    "Attribute 1 name",
    "Attribute 1 value(s)",
    "Attribute 1 visible",
    "Attribute 1 global",
    "Attribute 2 name",
    "Attribute 2 value(s)",
    "Attribute 2 visible",
    "Attribute 2 global",
]


def _decimal_str(value: Decimal | None) -> str:
    return "" if value is None else format(value, "f")


def _is_blocked(product: CanonicalProduct, validation: ValidationResult) -> bool:
    return product.source.row in validation.blocked_rows


def _base_row(product: CanonicalProduct, config: dict[str, Any]) -> dict[str, str | int]:
    woo = config.get("woocommerce", {})
    return {
        "Type": "simple",
        "SKU": product.sku or "",
        "GTIN, UPC, EAN, or ISBN": product.ean or "",
        "Name": product.name or "",
        "Published": int(woo.get("published", -1)),
        "Tax status": woo.get("tax_status", "taxable"),
        "Tax class": woo.get("tax_class", "standard"),
        "In stock?": 1 if (product.stock or 0) > 0 else 0,
        "Stock": "" if product.stock is None else product.stock,
        "Regular price": _decimal_str(product.price),
        "Categories": product.category or "",
        "Parent": "",
        "Attribute 1 name": "",
        "Attribute 1 value(s)": "",
        "Attribute 1 visible": "",
        "Attribute 1 global": "",
        "Attribute 2 name": "",
        "Attribute 2 value(s)": "",
        "Attribute 2 visible": "",
        "Attribute 2 global": "",
    }


def _attribute_pairs(product: CanonicalProduct, config: dict[str, Any]) -> list[tuple[str, str]]:
    woo = config.get("woocommerce", {})
    order = woo.get("attribute_order", ["color", "size"])
    labels = woo.get("attribute_labels", {"color": "Color", "size": "Size"})
    pairs: list[tuple[str, str]] = []
    for field in order:
        value = getattr(product, field, None)
        if value:
            pairs.append((str(labels.get(field, field.title())), str(value)))
    return pairs[:2]


def export_woocommerce_csv(products: list[CanonicalProduct], validation: ValidationResult, config: dict[str, Any], path: str | Path) -> int:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    exportable = [p for p in products if p.sku and not _is_blocked(p, validation)]

    simple = [p for p in exportable if not p.parent_sku]
    groups: dict[str, list[CanonicalProduct]] = defaultdict(list)
    for product in exportable:
        if product.parent_sku:
            groups[product.parent_sku].append(product)

    rows: list[dict[str, str | int]] = []
    for product in sorted(simple, key=lambda p: p.sku or ""):
        rows.append(_base_row(product, config))

    for parent_sku in sorted(groups):
        members = sorted(groups[parent_sku], key=lambda p: p.sku or "")
        first = members[0]
        parent = _base_row(first, config)
        parent.update({
            "Type": "variable",
            "SKU": parent_sku,
            "GTIN, UPC, EAN, or ISBN": "",
            "Name": first.parent_name or parent_sku,
            "In stock?": 1 if any((p.stock or 0) > 0 for p in members) else 0,
            "Stock": "",
            "Regular price": "",
            "Parent": "",
        })

        labels = config.get("woocommerce", {}).get("attribute_labels", {"color": "Color", "size": "Size"})
        order = config.get("woocommerce", {}).get("attribute_order", ["color", "size"])
        attr_slot = 1
        for field in order:
            values = sorted({str(getattr(p, field)) for p in members if getattr(p, field, None)})
            if not values or attr_slot > 2:
                continue
            parent[f"Attribute {attr_slot} name"] = str(labels.get(field, field.title()))
            parent[f"Attribute {attr_slot} value(s)"] = ", ".join(values)
            parent[f"Attribute {attr_slot} visible"] = 1
            parent[f"Attribute {attr_slot} global"] = 0
            attr_slot += 1
        rows.append(parent)

        for product in members:
            row = _base_row(product, config)
            row["Type"] = "variation"
            row["Parent"] = parent_sku
            row["Categories"] = ""
            for idx, (label, value) in enumerate(_attribute_pairs(product, config), start=1):
                row[f"Attribute {idx} name"] = label
                row[f"Attribute {idx} value(s)"] = value
                row[f"Attribute {idx} visible"] = 1
                row[f"Attribute {idx} global"] = 0
            rows.append(row)

    gtin_mode = config.get("woocommerce", {}).get("gtin_mode", "native")
    headers = list(CSV_HEADERS)
    if gtin_mode == "omit":
        headers.remove("GTIN, UPC, EAN, or ISBN")
        for row in rows:
            row.pop("GTIN, UPC, EAN, or ISBN", None)
    elif gtin_mode == "meta":
        idx = headers.index("GTIN, UPC, EAN, or ISBN")
        headers[idx] = "meta:_global_unique_id"
        for row in rows:
            row["meta:_global_unique_id"] = row.pop("GTIN, UPC, EAN, or ISBN", "")

    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="16324F")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(10, min(max_len + 2, 42))


def export_report_xlsx(
    path: str | Path,
    current_products: list[CanonicalProduct],
    validation: ValidationResult,
    changes: list[Change],
    unchanged: int,
    normalizations: list[NormalizationRecord],
    exported_rows: int,
) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "SUMMARY"

    by_severity = Counter(issue.severity for issue in validation.issues)
    change_counts = Counter(change.change_type for change in changes)
    unique_current_skus = {p.sku for p in current_products if p.sku}
    blocked_rows = validation.blocked_rows
    accepted_supplier_rows = len(current_products) - len(blocked_rows)
    synthesized_parent_rows = max(0, exported_rows - accepted_supplier_rows)
    summary_rows = [
        ["Metric", "Value"],
        ["Supplier rows", len(current_products)],
        ["Unique non-empty SKUs", len(unique_current_skus)],
        ["Blocked supplier rows", len(blocked_rows)],
        ["Accepted supplier rows", accepted_supplier_rows],
        ["Synthesized variable parent rows", synthesized_parent_rows],
        ["WooCommerce CSV rows", exported_rows],
        ["Unchanged products", unchanged],
        ["New products", change_counts.get("NEW_PRODUCT", 0)],
        ["Removed from supplier feed", change_counts.get("REMOVED_FROM_FEED", 0)],
        ["Price changes", change_counts.get("PRICE_CHANGED", 0)],
        ["Stock changes", change_counts.get("STOCK_CHANGED", 0)],
        ["Packaging changes", change_counts.get("PACKAGING_CHANGED", 0)],
        ["Errors", by_severity.get("ERROR", 0)],
        ["Warnings", by_severity.get("WARNING", 0)],
        ["Needs human review", by_severity.get("REVIEW", 0)],
        ["Normalization actions", len(normalizations)],
    ]
    for row in summary_rows:
        summary.append(row)
    _style_sheet(summary)

    changes_ws = wb.create_sheet("CHANGES")
    changes_ws.append(["SKU", "Change type", "Field", "Previous", "Current", "Previous row", "Current row", "Note"])
    for change in changes:
        changes_ws.append([change.sku, change.change_type, change.field, change.old_value, change.new_value, change.previous_row, change.current_row, change.note])
    _style_sheet(changes_ws)

    for severity in ("ERROR", "WARNING", "REVIEW"):
        ws = wb.create_sheet(severity + "S")
        ws.append(["Source file", "Sheet", "Row", "SKU", "Code", "Field", "Raw value", "Message"])
        for issue in validation.issues:
            if issue.severity == severity:
                ws.append([issue.source.file, issue.source.sheet, issue.source.row, issue.sku, issue.code, issue.field, str(issue.raw_value), issue.message])
        _style_sheet(ws)

    norms_ws = wb.create_sheet("NORMALIZATIONS")
    norms_ws.append(["Source file", "Sheet", "Row", "SKU", "Field", "Raw value", "Normalized value", "Rule"])
    for record in normalizations:
        norms_ws.append([record.source.file, record.source.sheet, record.source.row, record.sku, record.field, str(record.raw_value), str(record.normalized_value), record.rule])
    _style_sheet(norms_ws)

    wb.save(path)
