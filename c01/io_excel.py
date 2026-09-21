from __future__ import annotations

from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from .models import RawRow, SourceRef


def load_config(path: str | Path) -> dict[str, Any]:
    with open(path, "r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle)
    if not isinstance(config, dict):
        raise ValueError("Config must be a YAML mapping")
    return config


def read_supplier_xlsx(path: str | Path, config: dict[str, Any]) -> list[RawRow]:
    path = Path(path)
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet_name = config.get("sheet_name")
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} not found in {path.name}")
    sheet = workbook[sheet_name]

    header_row = int(config.get("header_row", 1))
    headers: dict[str, int] = {}
    for cell in sheet[header_row]:
        if cell.value is not None:
            headers[str(cell.value).strip()] = cell.column

    required = config.get("required_headers", [])
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"Missing required headers in {path.name}: {', '.join(missing)}")

    column_map: dict[str, str] = config["columns"]
    rows: list[RawRow] = []

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        values: dict[str, Any] = {}
        meta: dict[str, dict[str, Any]] = {}
        has_data = False
        for source_header, canonical_name in column_map.items():
            col_idx = headers.get(source_header)
            if col_idx is None:
                values[canonical_name] = None
                continue
            cell = sheet.cell(row=row_idx, column=col_idx)
            values[canonical_name] = cell.value
            meta[canonical_name] = {
                "data_type": cell.data_type,
                "number_format": cell.number_format,
            }
            if cell.value not in (None, ""):
                has_data = True

        if not has_data:
            continue

        rows.append(
            RawRow(
                source=SourceRef(file=path.name, sheet=sheet.title, row=row_idx),
                values=values,
                cell_meta=meta,
            )
        )

    workbook.close()
    return rows
